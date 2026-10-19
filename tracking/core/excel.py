from __future__ import annotations

import datetime
from dataclasses import dataclass
from typing import Any, Iterable

import openpyxl
from django.http import HttpResponse
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

XLSX_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
DATE_FORMAT = 'DD/MM/YYYY'
_BOLD = Font(bold=True)


@dataclass(frozen=True)
class ExcelColumn:
    key: str
    header: str
    width: float | None = None
    is_date: bool = False


@dataclass(frozen=True)
class ExcelSheet:
    title: str
    columns: list[ExcelColumn]
    rows: Iterable[dict[str, Any]]
    freeze: str | None = 'A2'
    total_label: str | None = None


def _cell(column: ExcelColumn, value: Any) -> Any:
    if column.is_date and value:
        return datetime.date.fromisoformat(str(value)[:10])
    return value


def _write_sheet(ws, sheet: ExcelSheet) -> None:
    ws.append([c.header for c in sheet.columns])
    for cell in ws[1]:
        cell.font = _BOLD
    rows = list(sheet.rows)
    for row in rows:
        ws.append([_cell(c, row.get(c.key)) for c in sheet.columns])
    if sheet.total_label:
        last = sheet.columns[-1].key
        ws.append([sheet.total_label, *([None] * (len(sheet.columns) - 2)), sum(r.get(last) or 0 for r in rows)])
        for cell in ws[ws.max_row]:
            cell.font = _BOLD
    for idx, column in enumerate(sheet.columns, start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = column.width if column.width is not None else max(len(column.header) + 2, 14)
        if column.is_date:
            for cell in ws[letter][1:]:
                cell.number_format = DATE_FORMAT
    if sheet.freeze:
        ws.freeze_panes = sheet.freeze


def excel_response(sheets: list[ExcelSheet], filename: str) -> HttpResponse:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet in sheets:
        _write_sheet(wb.create_sheet(sheet.title), sheet)
    response = HttpResponse(content_type=XLSX_CONTENT_TYPE)
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response
