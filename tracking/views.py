import datetime
import re
import openpyxl
from openpyxl.utils import get_column_letter
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render
from .detail_report import DETAIL_CONFIGS, EMPLOYEE_REPORT_CONFIG_DMS, EMPLOYEE_REPORT_CONFIG_OMS, INVENTORY_CONFIG, RECON_CONFIGS, DetailReportConfig, export_detail_report_rows, get_distinct_so_status, get_inventory_detail, get_inventory_diff_list, get_inventory_meta, run_detail_report
from .SystemsTracking import get_report

def export_detail_report_excel(config: DetailReportConfig, params: dict, filename: str) -> HttpResponse:
    rows = export_detail_report_rows(config, params)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Data'
    headers = [c.title or c.data_key for c in config.columns]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True)
    for row in rows:
        ws.append([row.get(c.data_key) for c in config.columns])
    for idx, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = max(len(header) + 2, 14)
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response

@login_required
def systems_tracking_page(request):
    return render(request, 'tracking/systems_tracking.html')

@login_required
def systems_tracking_report(request):
    return JsonResponse({'data': get_report()})

def recon_ecom_page(request):
    return render(request, 'tracking/recon_ecom.html')

def recon_ecom_report(request):
    return JsonResponse(run_detail_report(RECON_CONFIGS['ecom'], request.GET))

def recon_ecom_export(request):
    return export_detail_report_excel(RECON_CONFIGS['ecom'], request.GET, 'data_tracking_ecom.xlsx')

def recon_dms_duyet_page(request):
    return render(request, 'tracking/recon_dms_duyet.html')

def recon_dms_duyet_report(request):
    return JsonResponse(run_detail_report(RECON_CONFIGS['dms-duyet'], request.GET))

def recon_dms_duyet_export(request):
    return export_detail_report_excel(RECON_CONFIGS['dms-duyet'], request.GET, 'data_tracking_dms_duyet.xlsx')

def recon_dms_xuatkho_page(request):
    return render(request, 'tracking/recon_dms_xuatkho.html')

def recon_dms_xuatkho_report(request):
    return JsonResponse(run_detail_report(RECON_CONFIGS['dms-xuatkho'], request.GET))

def recon_dms_xuatkho_export(request):
    return export_detail_report_excel(RECON_CONFIGS['dms-xuatkho'], request.GET, 'data_tracking_dms_xuatkho.xlsx')

def recon_wrt_page(request):
    return render(request, 'tracking/recon_wrt.html')

def recon_wrt_report(request):
    return JsonResponse(run_detail_report(RECON_CONFIGS['wrt'], request.GET))

def recon_wrt_export(request):
    return export_detail_report_excel(RECON_CONFIGS['wrt'], request.GET, 'data_tracking_wrt.xlsx')

def recon_wrt_so_status_options(request):
    return JsonResponse({'options': get_distinct_so_status(RECON_CONFIGS['wrt'])})

def recon_inventory_page(request):
    return render(request, 'tracking/recon_inventory.html')

def recon_inventory_report(request):
    return JsonResponse(run_detail_report(INVENTORY_CONFIG, request.GET))

def recon_inventory_export(request):
    return export_detail_report_excel(INVENTORY_CONFIG, request.GET, 'data_tracking_inventory.xlsx')

def recon_inventory_meta(request):
    return JsonResponse(get_inventory_meta())

def recon_inventory_detail(request):
    item_code = (request.GET.get('item_code') or '').strip()
    whs_code = (request.GET.get('whs_code') or '').strip()
    status_item = (request.GET.get('status_item') or '').strip()
    if not item_code or not whs_code:
        return JsonResponse({'error': 'Thieu item_code hoac whs_code'}, status=400)
    return JsonResponse(get_inventory_detail(item_code, whs_code, status_item))

_DETAIL_EXPORT_SHEETS = (
    ('both', 'Hai bên có', (('doc_no', 'Số chứng từ', 24), ('wms_doc_no', 'Số chứng từ WMS', 24), ('sap_date', 'Ngày', 12), ('wms_date', 'Ngày WMS', 12), ('doc_type', 'Loại', 30), ('wms_doc_type', 'Loại WMS', 30), ('sap_qty', 'SAP', 10), ('wms_qty', 'WMS', 10), ('diff_qty', 'Lệch', 10))),
    ('sap_only', 'SAP có WMS không', (('doc_no', 'Số chứng từ', 24), ('doc_date', 'Ngày chứng từ', 14), ('doc_type', 'Loại', 30), ('qty', 'SL (+/-)', 10))),
    ('wms_only', 'WMS có SAP không', (('doc_no', 'Số chứng từ', 24), ('doc_date', 'Ngày chứng từ', 14), ('doc_type', 'Loại', 30), ('qty', 'SL (+/-)', 10))),
)

def _excel_cell(field, value):
    if value and field.endswith('_date'):
        return datetime.date.fromisoformat(str(value)[:10])
    return value

def _inventory_workbook(sheets, data, filename):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    bold = openpyxl.styles.Font(bold=True)
    for key, title, columns in sheets:
        ws = wb.create_sheet(title)
        ws.append([header for _, header, _ in columns])
        for cell in ws[1]:
            cell.font = bold
        rows = data[key]
        for row in rows:
            ws.append([_excel_cell(field, row.get(field)) for field, _, _ in columns])
        total_field = columns[-1][0]
        ws.append(['Tổng', *([None] * (len(columns) - 2)), sum((r[total_field] or 0 for r in rows))])
        for cell in ws[ws.max_row]:
            cell.font = bold
        for idx, (field, _, width) in enumerate(columns, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width
            if field.endswith('_date'):
                for cell in ws[get_column_letter(idx)][1:]:
                    cell.number_format = 'DD/MM/YYYY'
        ws.freeze_panes = 'A2'
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response

def recon_inventory_detail_export(request):
    item_code = (request.GET.get('item_code') or '').strip()
    whs_code = (request.GET.get('whs_code') or '').strip()
    status_item = (request.GET.get('status_item') or '').strip()
    if not item_code or not whs_code:
        return JsonResponse({'error': 'Thieu item_code hoac whs_code'}, status=400)
    safe = re.sub(r'[^A-Za-z0-9_-]', '', f'{item_code}_{whs_code}_{status_item}'.strip('_'))
    return _inventory_workbook(_DETAIL_EXPORT_SHEETS, get_inventory_detail(item_code, whs_code, status_item), f'chi_tiet_ton_kho_{safe}.xlsx')

def recon_inventory_diff_list(request):
    return JsonResponse(get_inventory_diff_list())

_DIFF_EXPORT_SHEETS = (
    ('both', 'Hai bên có', (('item_code', 'Mã hàng', 14), ('item_name', 'Tên hàng', 50), ('whs_code', 'Kho', 10), ('status_item', 'Trạng thái', 14), ('sap_qty', 'SAP', 12), ('wms_qty', 'WMS', 12), ('diff_qty', 'Lệch', 12))),
    ('sap_only', 'SAP có WMS không', (('item_code', 'Mã hàng', 14), ('item_name', 'Tên hàng', 50), ('whs_code', 'Kho', 10), ('status_item', 'Trạng thái', 14), ('sap_qty', 'SL', 12))),
    ('wms_only', 'WMS có SAP không', (('item_code', 'Mã hàng', 14), ('item_name', 'Tên hàng', 50), ('whs_code', 'Kho', 10), ('status_item', 'Trạng thái', 14), ('wms_qty', 'SL', 12))),
)

def recon_inventory_diff_export(request):
    return _inventory_workbook(_DIFF_EXPORT_SHEETS, get_inventory_diff_list(), 'ton_kho_lech_sap_wms.xlsx')

def detail_ecom_page(request):
    return render(request, 'tracking/detail_ecom.html')

def detail_ecom_report(request):
    return JsonResponse(run_detail_report(DETAIL_CONFIGS['ecom'], request.GET))

def detail_ecom_export(request):
    return export_detail_report_excel(DETAIL_CONFIGS['ecom'], request.GET, 'chi_tiet_ecom.xlsx')

def detail_dms_duyet_page(request):
    return render(request, 'tracking/detail_dms_duyet.html')

def detail_dms_duyet_report(request):
    return JsonResponse(run_detail_report(DETAIL_CONFIGS['dms-duyet'], request.GET))

def detail_dms_duyet_export(request):
    return export_detail_report_excel(DETAIL_CONFIGS['dms-duyet'], request.GET, 'chi_tiet_dms_duyet.xlsx')

def detail_dms_xuatkho_page(request):
    return render(request, 'tracking/detail_dms_xuatkho.html')

def detail_dms_xuatkho_report(request):
    return JsonResponse(run_detail_report(DETAIL_CONFIGS['dms-xuatkho'], request.GET))

def detail_dms_xuatkho_export(request):
    return export_detail_report_excel(DETAIL_CONFIGS['dms-xuatkho'], request.GET, 'chi_tiet_dms_xuatkho.xlsx')

def detail_wrt_page(request):
    return render(request, 'tracking/detail_wrt.html')

def detail_wrt_report(request):
    return JsonResponse(run_detail_report(DETAIL_CONFIGS['wrt'], request.GET))

def detail_wrt_export(request):
    return export_detail_report_excel(DETAIL_CONFIGS['wrt'], request.GET, 'chi_tiet_wrt.xlsx')

def detail_employee_dms_page(request):
    return render(request, 'tracking/detail_employee_dms.html')

def detail_employee_dms_report(request):
    return JsonResponse(run_detail_report(EMPLOYEE_REPORT_CONFIG_DMS, request.GET))

def detail_employee_dms_export(request):
    return export_detail_report_excel(EMPLOYEE_REPORT_CONFIG_DMS, request.GET, 'chi_tiet_don_hang_theo_nhan_vien_dms.xlsx')

def detail_employee_oms_page(request):
    return render(request, 'tracking/detail_employee_oms.html')

def detail_employee_oms_report(request):
    return JsonResponse(run_detail_report(EMPLOYEE_REPORT_CONFIG_OMS, request.GET))

def detail_employee_oms_export(request):
    return export_detail_report_excel(EMPLOYEE_REPORT_CONFIG_OMS, request.GET, 'chi_tiet_don_hang_theo_nhan_vien_oms.xlsx')
