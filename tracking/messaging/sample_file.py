from __future__ import annotations

from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

SHEET_TITLE = 'Danh sách khách hàng'
SAMPLE_FILENAME = 'mau_danh_sach_gui_tin.xlsx'
HEADERS = ('Tên khách hàng', 'Số điện thoại', 'Trạng thái gửi')
SAMPLE_ROWS = (('Nguyễn Văn A', '0879573727', None), ('Trần Thị B', '0888236776', None))
PHONE_COL = 2

CONFIG_SHEET_TITLE = 'Cấu hình mẫu tin'
CONFIG_FILENAME = 'mau_cau_hinh_mau_tin.xlsx'
INFO_HEADERS = ('Thông tin', 'Giá trị')
INFO_ROWS = (
    ('Mã', '123456'),
    ('Tên', 'ZNS Khuyến mãi PGI'),
    ('Chủ hàng', 'PGI'),
    ('Nội dung', 'Chào <name>, PGI gửi tặng bạn voucher giảm <giam> mã <voucher_code>, áp dụng từ <start_date> đến <expired_date>.'),
)
VAR_HEADERS = ('Tên biến', 'Kiểu dữ liệu', 'Dữ liệu', 'Bắt buộc', 'Theo KH')
VAR_ROWS = (
    ('name', 'Text', None, 'Có', 'Có'),
    ('giam', 'Text', '10%', 'Có', 'Không'),
    ('voucher_code', 'Text', 'PGI10', 'Có', 'Không'),
    ('start_date', 'Date', '01/11/2026', 'Có', 'Không'),
    ('expired_date', 'Datetime', '22:00 30/11/2026', 'Có', 'Không'),
)
VAR_INPUT_ROWS = 50
TYPE_CHOICES = '"Text,Number,Date,Datetime"'
YES_NO_CHOICES = '"Có,Không"'
NOTE = 'Ghi chú: Theo KH = Có thì Dữ liệu để trống, mỗi khách lấy giá trị từ cột cùng tên trong file danh sách khách hàng.'

_LINE = Side(style='thin', color='D9DDE4')
_BORDER = Border(left=_LINE, right=_LINE, top=_LINE, bottom=_LINE)
_HEAD_FILL = PatternFill('solid', fgColor='F1F3F7')
_HEAD_FONT = Font(bold=True, color='374151')


def _table(ws: Worksheet, start_row: int, headers: tuple[str, ...], rows: tuple[tuple[Any, ...], ...], text_cols: tuple[int, ...] = ()) -> None:
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=header)
        cell.font = _HEAD_FONT
        cell.fill = _HEAD_FILL
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[start_row].height = 24
    for offset, values in enumerate(rows, start=1):
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=start_row + offset, column=col, value=value)
            cell.border = _BORDER
            if col in text_cols:
                cell.number_format = '@'


def sample_workbook() -> tuple[openpyxl.Workbook, str]:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_TITLE
    _table(ws, 1, HEADERS, SAMPLE_ROWS, (PHONE_COL,))
    for col in range(1, len(HEADERS) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 22
    ws.freeze_panes = 'A2'
    return wb, SAMPLE_FILENAME


def template_sample_workbook() -> tuple[openpyxl.Workbook, str]:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = CONFIG_SHEET_TITLE
    _table(ws, 1, INFO_HEADERS, INFO_ROWS, (2,))
    for offset, (key, _) in enumerate(INFO_ROWS, start=2):
        ws.cell(row=offset, column=2).alignment = Alignment(wrap_text=True, vertical='top')
        if key == 'Nội dung':
            ws.merge_cells(start_row=offset, start_column=2, end_row=offset, end_column=len(VAR_HEADERS))
            ws.row_dimensions[offset].height = 48
    ws.cell(row=2 + INFO_ROWS.index(next(r for r in INFO_ROWS if r[0] == 'Chủ hàng')), column=2).number_format = '@'

    var_row = len(INFO_ROWS) + 3
    padding = tuple((None,) * len(VAR_HEADERS) for _ in range(VAR_INPUT_ROWS - len(VAR_ROWS)))
    _table(ws, var_row, VAR_HEADERS, VAR_ROWS + padding, (3,))
    first, last = var_row + 1, var_row + VAR_INPUT_ROWS
    types = DataValidation(type='list', formula1=TYPE_CHOICES, allow_blank=True)
    yes_no = DataValidation(type='list', formula1=YES_NO_CHOICES, allow_blank=True)
    ws.add_data_validation(types)
    ws.add_data_validation(yes_no)
    types.add(f'B{first}:B{last}')
    yes_no.add(f'D{first}:E{last}')
    owners = DataValidation(type='list', formula1='"PGI,AWG"', allow_blank=False)
    ws.add_data_validation(owners)
    owners.add('B4')
    ws.cell(row=last + 2, column=1, value=NOTE).font = Font(italic=True, color='6B7280')

    for col, width in enumerate((22, 16, 24, 12, 12), start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    return wb, CONFIG_FILENAME
